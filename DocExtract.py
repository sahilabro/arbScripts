from docx import Document
import docx
from docx import Document
from docx.document import Document as _Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph


def iter_block_items(parent):
    """
    Generate a reference to each paragraph and table child within *parent*,
    in document order. Each returned value is an instance of either Table or
    Paragraph. *parent* would most commonly be a reference to a main
    Document object, but also works for a _Cell object, which itself can
    contain paragraphs and tables.
    """
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
        # print(parent_elm.xml)
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("something's not right")

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


document = Document('IAG Test2 Infrastructure Design v0.4.docx')
iters = iter_block_items(document)
a = 0
b = 0
i= 0
host = ''
masterlist = []
for block in iters:
    if isinstance(block,Paragraph):
        if 'The table below defines the properties of the host.' in block.text:
            b = 1
        if 'Anti-Affinity Hosts' in block.text:
            a = 1
        elif 'No Anti-Affinity' in block.text:
            a = 0
            continue
    if b== 1 and isinstance(block,Table):
        #print('Host is ' + block.cell(0,1).text)
        host = block.cell(0,1).text
        b = 0
    if a == 1 and isinstance(block,Table):        
        #print([x.text for x in [cell for cell in block.column_cells(0)]][1::])
        forbiddenfruit = [x.text for x in [cell for cell in block.column_cells(0)]][1::]
        forbiddenfruit.append(host)
        masterlist.append(forbiddenfruit)
        i += 1
        a = 0

    #print(block.text if isinstance(block, Paragraph) else '<table>')
print(masterlist)


